"""Read-only loading of the CorroborAI extracts.

Source files are never written to. Everything is read with openpyxl/pandas and
kept in memory.
"""

from __future__ import annotations

import datetime as dt
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
import pandas as pd

SOURCE_FILE = "Employe_Source_Anonymise_VF.xlsx"
TARGET_FILE = "Employe_Destination_Anonymise_VF.xlsx"
MAPPING_FILE = "Mapping.xlsx"
MOTIF_FILE = "Motif de la situation d'emploi.xlsx"
DETAIL_FILE = "détail_du_poste.xlsx"

EXCEL_EPOCH = dt.date(1899, 12, 30)


@dataclass
class Inputs:
    source: pd.DataFrame  # Système A - RH
    target: pd.DataFrame  # Système B - Temps
    motif: dict  # CodeCatégorieStatut -> {"remphor": ..., "acces": ...}
    detail: pd.DataFrame  # historique du détail de poste, une ligne par effet
    mapping_rows: dict = field(default_factory=dict)  # champ source -> ligne Excel dans Mapping.xlsx


def fix_mojibake(value):
    """Repair UTF-8 text that was decoded as latin-1 ("Absence complÃ¨te")."""
    if isinstance(value, str):
        try:
            return value.encode("latin-1").decode("utf-8")
        except (UnicodeEncodeError, UnicodeDecodeError):
            return value
    return value


def _load_motif(path: Path) -> dict:
    df = pd.read_excel(path)
    out = {}
    for _, row in df.iterrows():
        out[int(row["CodeCatégorieStatut"])] = {
            "remphor": int(row["CodeStatutSystèmeExterne"]),
            "acces": int(row["CodeGestionAccès"]),
        }
    return out


def _load_detail(path: Path) -> pd.DataFrame:
    """The job-detail extract arrives as one comma-separated text column."""
    ws = openpyxl.load_workbook(path, read_only=True, data_only=True).active
    header = None
    records = []
    for raw in ws.iter_rows(values_only=True):
        if raw[0] is None:
            continue
        text = str(raw[0])
        if header is None:
            header = [h.strip() for h in text.split(",")]
            continue
        records.append(dict(zip(header, text.split(","))))
    df = pd.DataFrame(records)
    df["IdentifiantPoste"] = df["IdentifiantPoste"].astype(int)
    df["DateEffetAffectation"] = df["DateEffetAffectation"].astype(int)
    df["HeuresSemaineContrat"] = df["HeuresSemaineContrat"].astype(float)
    df["HeuresJourContrat"] = df["HeuresJourContrat"].astype(float)
    df["CodeDirectionAffectée"] = df["CodeDirectionAffectée"].astype(int)
    df["date_effet"] = df["DateEffetAffectation"].map(lambda v: EXCEL_EPOCH + dt.timedelta(days=v))
    return df


def _mapping_rows(path: Path) -> dict:
    """Index of each source field in Mapping.xlsx, for traceable citations."""
    ws = openpyxl.load_workbook(path, read_only=True, data_only=True)["Mapping"]
    found = {}
    for idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
        if len(row) > 1 and row[1]:
            for part in str(row[1]).split("\n"):
                key = part.strip()
                found.setdefault(key, idx)
    return found


def load_inputs(root: str | Path) -> Inputs:
    root = Path(root)
    source = pd.read_excel(root / SOURCE_FILE)
    target = pd.read_excel(root / TARGET_FILE)
    return Inputs(
        source=source,
        target=target,
        motif=_load_motif(root / MOTIF_FILE),
        detail=_load_detail(root / DETAIL_FILE),
        mapping_rows=_mapping_rows(root / MAPPING_FILE),
    )
