"""Perturbation tests: inject known errors into a copy of the extracts and check they are caught.

This proves the rules do not depend on the specific records in the official set.
"""

import os
import shutil
from pathlib import Path

import openpyxl
import pytest

from corroborai.engine import run_engine
from corroborai.loader import load_inputs

DATA = Path(os.environ.get("CORROBORAI_DATA", Path(__file__).resolve().parents[2] / "corroborai-participants"))
pytestmark = pytest.mark.skipif(not (DATA / "Employe_Source_Anonymise_VF.xlsx").exists(), reason="extracts not available")
XLSX = ["Employe_Source_Anonymise_VF.xlsx", "Employe_Destination_Anonymise_VF.xlsx", "Mapping.xlsx",
        "Motif de la situation d'emploi.xlsx", "détail_du_poste.xlsx"]


@pytest.fixture
def copy(tmp_path):
    for name in XLSX + ["Mapping.xlsx"]:
        shutil.copy(DATA / name, tmp_path / name)
    return tmp_path


def _target_cell(path, person, column_name, nth=0):
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    col = headers.index(column_name) + 1
    rows = [r for r in range(2, ws.max_row + 1) if ws.cell(r, headers.index("personId") + 1).value == person]
    return wb, ws, ws.cell(rows[nth], col), rows


def _verdict(root, matricule, type_aff):
    res, _ = run_engine(load_inputs(root))
    hits = [r for r in res if str(r.matricule) == str(matricule) and r.type_affectation == type_aff]
    assert hits, (matricule, type_aff)
    return hits[0]


def test_injected_contract_error_is_caught(copy):
    path = copy / "Employe_Destination_Anonymise_VF.xlsx"
    wb, ws, cell, _ = _target_cell(path, 8142123, "contractTypeCode")
    assert cell.value == "JWN"
    cell.value = "KELH"  # wrong code for a permanent full-time "V" employee
    wb.save(path)
    row = _verdict(copy, 8142123, "P")
    assert row.verdict == "ANOMALIE"
    assert any(f.rule_id == "R15" and f.status == "ECART" for f in row.findings)


def test_injected_missing_target_row_is_caught(copy):
    path = copy / "Employe_Destination_Anonymise_VF.xlsx"
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    headers = [c.value for c in ws[1]]
    pid = headers.index("personId") + 1
    for r in range(ws.max_row, 1, -1):
        if ws.cell(r, pid).value == 8644330:
            ws.delete_rows(r)
    wb.save(path)
    row = _verdict(copy, 8644330, "P")
    assert row.target_index is None
    assert row.verdict == "ANOMALIE"


def test_injected_unknown_person_in_target(copy):
    path = copy / "Employe_Destination_Anonymise_VF.xlsx"
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    last = ws.max_row
    for c in range(1, ws.max_column + 1):
        ws.cell(last + 1, c).value = ws.cell(last, c).value
    headers = [c.value for c in ws[1]]
    ws.cell(last + 1, headers.index("personId") + 1).value = 1234567
    wb.save(path)
    res, meta = run_engine(load_inputs(copy))
    orphans = [r for r in res if r.source_index is None]
    assert len(orphans) == 1 and orphans[0].verdict == "ANOMALIE"


def test_missing_motif_entry_goes_to_arbitration(copy):
    path = copy / "Motif de la situation d'emploi.xlsx"
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    for r in range(ws.max_row, 1, -1):
        if ws.cell(r, 1).value == 807:
            ws.delete_rows(r)
    wb.save(path)
    row = _verdict(copy, 7603160, "P")
    finding = next(f for f in row.findings if f.rule_id == "R24")
    assert finding.status == "A_ARBITRER"
    assert row.verdict == "À ARBITRER"
